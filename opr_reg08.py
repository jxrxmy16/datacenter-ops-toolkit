#!/usr/bin/env python3
"""
opr_reg08.py – Rellena el formulario OPR-REG-08 (ingreso / instalación de equipos)
a partir de la planilla de equipos que descarga el cliente.

Usage:
    python opr_reg08.py equipos.xlsx --cliente SOLUTIONS --ticket SD5229372
    python opr_reg08.py equipos.csv  --cliente SOLUTIONS --ticket SD5229372 --fecha "03/10/2026 12:00 PM"
    python opr_reg08.py equipos.xlsx --plantilla plantillas/OPR-REG-08.docx -o salida.docx
    python opr_reg08.py equipos.xlsx --columnas      # solo muestra cómo se detectaron las columnas

La planilla necesita una fila por equipo. Los encabezados se reconocen por nombre
(Equipamiento, Marca, Modelo, Consumo/kW, N° de serie, Sala, Rack, U...), en
cualquier orden y aunque haya filas de título antes. Las filas de equipos del
Word se agregan o quitan según la cantidad de equipos.
"""

import argparse
import copy
import csv
import re
import sys
import unicodedata
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.table import _Cell
from rich.console import Console
from rich.table import Table

console = Console()

DEFAULT_TEMPLATE = Path(__file__).parent / "plantillas" / "OPR-REG-08.docx"

# Campo interno -> encabezados aceptados (ya normalizados: minúsculas, sin tildes ni signos)
ALIASES = {
    "equipamiento": {"equipamiento", "equipo", "tipo", "tipo de equipo", "tipo equipo", "descripcion", "familia"},
    "marca":        {"marca", "fabricante", "brand", "vendor"},
    "modelo":       {"modelo", "model"},
    "kw":           {"consumo", "kw", "consumo kw", "consumo kw declarado", "kw declarado", "potencia",
                     "potencia kw", "consumo declarado"},
    "serie":        {"serie", "n de serie", "numero de serie", "nro de serie", "nro serie", "n serie",
                     "serial", "serial number", "sn", "s n"},
    "sala":         {"sala"},
    "rack":         {"rack"},
    "u":            {"u", "u rack", "us", "unidad", "unidades", "posicion", "posicion rack", "u inicial"},
    "rfid":         {"rfid"},
    "cliente":      {"cliente", "empresa", "razon social"},
    "ticket":       {"ticket", "n ticket", "nro ticket", "solicitud", "sd"},
}

# Índice de <w:tc> dentro de cada fila de equipo (las celdas combinadas cuentan una vez)
ARMADO_COLS = {"num": 0, "equipamiento": 1, "marca": 2, "modelo": 3, "kw": 4, "serie": 6, "rfid": 7}
BLANCA_COLS = {"num": 0, "sala": 1, "rack": 2, "equipamiento": 3, "marca": 4, "modelo": 5, "kw": 6,
               "serie": 8, "rfid": 9, "u": 10}
HEADER_COLS = {"cliente": 1, "ticket": 3, "fecha": 5}   # fila 0 de ambas tablas
FIRST_EQ_ROW = 5                                        # filas 1-4 son el encabezado combinado


# ── lectura de la planilla ────────────────────────────────────────────────────

def normalize(text) -> str:
    text = unicodedata.normalize("NFKD", str(text or "")).encode("ascii", "ignore").decode()
    text = re.sub(r"[^a-z0-9]+", " ", text.lower())
    return text.strip()


def match_field(header: str):
    h = normalize(header)
    for field, names in ALIASES.items():
        if h in names:
            return field
    for field, names in ALIASES.items():
        if any(len(n) > 3 and h.startswith(n) for n in names):
            return field
    return None


def format_value(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float):
        value = f"{value:.3f}".rstrip("0").rstrip(".").replace(".", ",")
    return str(value).strip()


def read_rows(path: Path) -> list[list]:
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        ws = load_workbook(path, data_only=True, read_only=True).active
        return [list(r) for r in ws.iter_rows(values_only=True)]
    raw = path.read_bytes()
    text = raw.decode("utf-8-sig") if raw.startswith(b"\xef\xbb\xbf") else raw.decode("utf-8", errors="replace")
    if "�" in text:
        text = raw.decode("latin-1")
    dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    return list(csv.reader(text.splitlines(), dialect))


def load_equipment(path: Path):
    rows = read_rows(path)
    for idx, row in enumerate(rows[:30]):
        mapping = {}
        for col, header in enumerate(row):
            field = match_field(header)
            if field and field not in mapping:
                mapping[field] = col
        if len({"equipamiento", "marca", "modelo", "serie"} & mapping.keys()) >= 2:
            break
    else:
        sys.exit(f"[error] No se encontró una fila de encabezados reconocible en {path}")

    equipos = []
    for row in rows[idx + 1:]:
        eq = {f: format_value(row[c]) if c < len(row) else "" for f, c in mapping.items()}
        if any(eq.get(f) for f in ("equipamiento", "marca", "modelo", "serie")):
            equipos.append(eq)
    return rows[idx], mapping, equipos


# ── escritura del Word ────────────────────────────────────────────────────────

def set_cell(tc, table, text: str):
    """Escribe texto en la celda conservando el formato del primer párrafo."""
    cell = _Cell(tc, table)
    para = cell.paragraphs[0]
    for extra in cell.paragraphs[1:]:
        extra._p.getparent().remove(extra._p)
    if para.runs:
        para.runs[0].text = text
        for run in para.runs[1:]:
            run._r.getparent().remove(run._r)
    else:
        run = para.add_run(text)
        # Copia el formato de la marca de párrafo (fuente/tamaño definidos en la plantilla)
        mark = para._p.pPr.find(qn("w:rPr")) if para._p.pPr is not None else None
        if mark is not None:
            run._r.insert(0, copy.deepcopy(mark))


def resize_equipment_rows(table, count: int):
    """Deja exactamente `count` filas de equipo clonando o quitando la última."""
    eq_rows = [r._tr for r in table.rows[FIRST_EQ_ROW:] if re.fullmatch(
        r"eq\d+", normalize("".join(_Cell(r._tr.tc_lst[0], table).text.split())))]
    last = eq_rows[-1]
    for _ in range(count - len(eq_rows)):
        new = copy.deepcopy(last)
        last.addnext(new)
        last = new
        eq_rows.append(new)
    for tr in eq_rows[count:]:
        tr.getparent().remove(tr)
    return eq_rows[:count]


def fill_table(table, cols: dict, equipos: list, header: dict, prefix: str):
    header_tcs = table.rows[0]._tr.tc_lst
    for field, i in HEADER_COLS.items():
        set_cell(header_tcs[i], table, header[field])
    for n, (tr, eq) in enumerate(zip(resize_equipment_rows(table, len(equipos)), equipos), 1):
        tcs = tr.tc_lst
        values = {**eq, "num": f"{prefix}{n}"}
        for field, i in cols.items():
            set_cell(tcs[i], table, values.get(field, ""))


def build_document(template: Path, output: Path, equipos: list, header: dict):
    doc = Document(template)
    if len(doc.tables) < 2:
        sys.exit(f"[error] La plantilla {template} no tiene las tablas esperadas")
    fill_table(doc.tables[0], ARMADO_COLS, equipos, header, "Eq")
    fill_table(doc.tables[1], BLANCA_COLS, equipos, header, "EQ")
    doc.save(output)


# ── CLI ───────────────────────────────────────────────────────────────────────

def main():
    ap = argparse.ArgumentParser(description="Rellena el OPR-REG-08 desde la planilla de equipos del cliente")
    ap.add_argument("planilla", type=Path, help="Archivo .xlsx o .csv descargado de la plataforma")
    ap.add_argument("--cliente", help="Nombre del cliente (si no viene en la planilla)")
    ap.add_argument("--ticket", help="N° de ticket, también se usa como RFID (si no viene en la planilla)")
    ap.add_argument("--fecha", help='Fecha/hora (default: ahora, ej. "03/10/2026 12:00 PM")')
    ap.add_argument("--plantilla", type=Path, default=DEFAULT_TEMPLATE, help="Plantilla OPR-REG-08 en blanco")
    ap.add_argument("-o", "--output", type=Path, help="Archivo .docx de salida")
    ap.add_argument("--columnas", action="store_true", help="Solo muestra las columnas detectadas y sale")
    args = ap.parse_args()

    if not args.planilla.exists():
        sys.exit(f"[error] No existe {args.planilla}")
    headers, mapping, equipos = load_equipment(args.planilla)

    t = Table(title=f"Columnas detectadas en {args.planilla.name}")
    t.add_column("Encabezado en planilla")
    t.add_column("Campo")
    for col, h in enumerate(headers):
        field = next((f for f, c in mapping.items() if c == col), None)
        t.add_row(str(h or ""), field or "[dim]ignorada[/dim]")
    console.print(t)
    if args.columnas:
        return

    first = equipos[0] if equipos else {}
    cliente = args.cliente or first.get("cliente", "")
    ticket = args.ticket or first.get("ticket", "")
    if not equipos:
        sys.exit("[error] La planilla no tiene equipos")
    if not cliente or not ticket:
        sys.exit("[error] Falta --cliente o --ticket (no vienen en la planilla)")
    for eq in equipos:
        eq["rfid"] = eq.get("rfid") or ticket

    if not args.plantilla.exists():
        sys.exit(f"[error] No existe la plantilla {args.plantilla} (usa --plantilla)")
    header = {
        "cliente": cliente,
        "ticket": ticket,
        "fecha": args.fecha or datetime.now().strftime("%d/%m/%Y %I:%M %p"),
    }
    output = args.output or Path(f"OPR-REG-08_Ingreso_Equipos_{normalize(cliente).replace(' ', '_').upper()}_{ticket}.docx")
    build_document(args.plantilla, output, equipos, header)

    missing = [f for f in ("sala", "rack", "u", "kw") if f not in mapping]
    console.print(f"[green]✓[/green] {len(equipos)} equipos → [bold]{output}[/bold]")
    if missing:
        console.print(f"[yellow]Aviso:[/yellow] la planilla no trae {', '.join(missing)}; esas celdas quedan vacías")


if __name__ == "__main__":
    main()
